# -*- coding: utf-8 -*-
"""xbdoc 纯函数测试：分词/切片/BM25/注入截断。零第三方依赖。

跑法（插件目录下任选其一）：
    python tests/test_retrieval.py
    python -m pytest tests/ -q
"""
import math
import os
import sys
from collections import Counter

sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.abspath(__file__))))

from xbdoc_inject import build_system_text, build_workspace_text, truncate_text
from xbdoc_retrieval import chunk_text, score_chunk_bm25, tokenize


def _idf(counters):
    n = len(counters)
    df = Counter()
    for c in counters:
        for t in c.keys():
            df[t] += 1
    return n, sum(sum(c.values()) for c in counters) / n, {
        t: math.log((n - f + 0.5) / (f + 0.5) + 1.0) for t, f in df.items()
    }


def test_chunk_long_single_paragraph():
    # 回归：超长单段曾经被整体丢弃（0 切片）
    assert len(chunk_text("a" * 2000, 1500, 200)) == 2
    assert len(chunk_text("x" * 5000, 1500, 200)) == 4
    assert len(chunk_text("line1\nline2\n\nline3", 1500, 200)) == 1
    assert chunk_text("", 1500, 200) == []


def test_bm25_rare_term_wins():
    docs = [
        "apple apple apple banana",
        "banana " + " ".join(f"w{i}" for i in range(50)),
        "cherry cherry",
    ]
    counters = [Counter(tokenize(d)) for d in docs]
    n, avg, idf = _idf(counters)
    scores = [score_chunk_bm25(tokenize("apple"), c, sum(c.values()), avg, idf)
              for c in counters]
    assert scores[0] > scores[1] and scores[0] > scores[2], scores


def test_bm25_length_norm():
    short_c = Counter(tokenize("apple"))
    long_c = Counter(tokenize("apple " + " ".join(f"u{i}" for i in range(100))))
    n, avg, idf = _idf([short_c, long_c])
    s_short = score_chunk_bm25(tokenize("apple"), short_c, sum(short_c.values()), avg, idf)
    s_long = score_chunk_bm25(tokenize("apple"), long_c, sum(long_c.values()), avg, idf)
    assert s_short > s_long > 0, (s_short, s_long)


def test_bm25_edges():
    c = Counter(tokenize("apple"))
    n, avg, idf = _idf([c])
    assert score_chunk_bm25([], c, 1, avg, idf) == 0.0
    assert score_chunk_bm25(tokenize("apple"), Counter(), 0, 0.0, {}) == 0.0
    assert score_chunk_bm25(tokenize("zzz"), c, 1, avg, idf) == 0.0


def test_tokenize_cjk_ext():
    # 日文假名/韩文/全角不再是零 token（旧正则只覆盖基本汉字块）；用码点构造，杜绝文件编码干扰
    s = "".join(chr(c) for c in (
        0x3053, 0x3093, 0x306B, 0x3061, 0x306F,  # こんにちは
        0x4E16, 0x754C,  # 世界
        0xD55C, 0xAE00,  # 한글
        0xFF21, 0xFF22,  # ＡＢ全角
    ))
    toks = tokenize(s)
    # 注意全角大写会被 lower() 转为全角小写（查询侧同样处理，对称可匹配）
    assert toks == [chr(c) for c in (
        0x3053, 0x3093, 0x306B, 0x3061, 0x306F, 0x4E16, 0x754C, 0xD55C, 0xAE00, 0xFF41, 0xFF42)], \
        [hex(ord(t)) for t in toks]
    assert tokenize("hello world") == ["hello", "world"]


def test_docx_tables():
    try:
        import docx
    except ImportError:
        return  # 环境无 python-docx 时跳过（AstrBot 侧安装后生效）
    import io
    from xbdoc_retrieval import extract_text_from_bytes
    doc = docx.Document()
    doc.add_paragraph("正文段落")
    t = doc.add_table(rows=1, cols=2)
    t.cell(0, 0).text = "表左"
    t.cell(0, 1).text = "表右"
    buf = io.BytesIO()
    doc.save(buf)
    text = extract_text_from_bytes(".docx", buf.getvalue())
    assert "正文段落" in text and "表左" in text and "表右" in text, text


def test_truncate():
    assert truncate_text("abc", 10) == "abc"
    # 过小上限：总长（含截断标记）严格 ≤ max_chars，不再超长
    out_small = truncate_text("a" * 100, 10)
    assert len(out_small) == 10 and out_small.endswith("…(截断)"), repr(out_small)
    out = truncate_text("a" * 1000, 600)
    assert out.endswith("…(截断)") and len(out) == 600, len(out)
    assert truncate_text("a" * 9000, 0) == "a" * 9000  # 0=不限制，完整注入
    assert truncate_text("a" * 9000, -5) == "a" * 9000
    # 专属提示词永不截断：build 段階只截文档部分
    out = build_system_text(["a" * 1000], "PROMPT", 600)
    assert "PROMPT" in out and out.endswith("PROMPT")
    ws = build_workspace_text([("f.md", "b" * 100)], "P", 1000)
    assert "/workspace/f.md" in ws and ws.endswith("P")
    # workspace 挂载头计入预算：含头总长不超 max_chars
    ws_tight = build_workspace_text([("f.md", "b" * 500)], "", 100)
    assert len(ws_tight) <= 100, len(ws_tight)


if __name__ == "__main__":
    test_chunk_long_single_paragraph()
    test_bm25_rare_term_wins()
    test_bm25_length_norm()
    test_bm25_edges()
    test_tokenize_cjk_ext()
    test_docx_tables()
    test_truncate()
    print("test_retrieval PASSED")
