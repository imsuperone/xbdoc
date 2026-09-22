"""xbdoc 纯函数：分词、切片、检索计分与文本提取。

无 AstrBot 依赖，可独立测试。
"""

import io
import re
from collections import Counter
from typing import List


# ======================================================================
# 纯函数：分词、切片、检索计分与文本提取
# ======================================================================

_TOKEN_RE = re.compile(
    r"[A-Za-z0-9_]+"
    r"|[\u1100-\u11ff\u3040-\u30ff\u3130-\u318f\u3400-\u4dbf\u4e00-\u9fff"
    r"\uac00-\ud7af\uf900-\ufaff\uff00-\uffef\U00020000-\U0002ebef]",
    re.UNICODE,
)
ALLOWED_SUFFIXES = {
    ".md", ".markdown", ".txt", ".json", ".csv", ".log",
    ".yaml", ".yml", ".toml", ".ini", ".cfg",
    ".html", ".htm", ".pdf", ".docx",
}

# 单次提取上限：超大 PDF 全量解码会吃爆内存/切片，切片 JSON 也会失控；注入侧另有上限
MAX_EXTRACT_CHARS = 500_000


def _cap_extract(text: str) -> str:
    if len(text) > MAX_EXTRACT_CHARS:
        return text[:MAX_EXTRACT_CHARS] + "\n…(原文超长截断)"
    return text


def tokenize(text: str) -> List[str]:
    """多语言混合分词：英文按词，中文/假名/谚文/扩展汉字按单字，统一小写。"""
    return [t.lower() for t in _TOKEN_RE.findall(text or "")]


def _overlap_tail(buf: str, overlap: int) -> str:
    """取 overlap 尾巴时尽量不断开英文单词：有空格时从空格后起，无空格（纯中文）保持字符级。"""
    tail = buf[-overlap:]
    # 只在尾巴较大且存在较早的断点时才 snap，避免 overlap 被削到过小
    for sep in ("\n", " ", "\t"):
        i = tail.find(sep)
        if 0 <= i < len(tail) - 20:
            return tail[i + 1:]
    return tail


def chunk_text(text: str, chunk_size: int = 1500, overlap: int = 200) -> List[str]:
    """段落优先滑动切片，保持上下文连贯与切片上限。"""
    text = (text or "").strip()
    if not text:
        return []
    chunk_size = max(200, int(chunk_size or 1500))
    overlap = max(0, min(int(overlap or 0), chunk_size - 50))

    paras = re.split(r"\n\s*\n|\n(?=#{1,6}\s)|\n", text)
    chunks: List[str] = []
    buf = ""

    for p in (p.strip() for p in paras if p.strip()):
        candidate = f"{buf}\n{p}".strip() if buf else p
        if len(candidate) <= chunk_size:
            buf = candidate
        else:
            if buf:
                chunks.append(buf)
                buf = (_overlap_tail(buf, overlap) + "\n" + p).strip() if overlap else p
            else:
                # buf 为空但单段已超长：不能丢弃，直接以该段为起点再硬切
                buf = p
            while len(buf) > chunk_size:
                chunks.append(buf[:chunk_size])
                buf = buf[chunk_size - overlap:].strip() if overlap else buf[chunk_size:].strip()
    if buf:
        chunks.append(buf)
    return [c for c in chunks if c.strip()]


def score_chunk_bm25(
    query_tokens: List[str],
    tf: Counter,
    doc_len: int,
    avg_len: float,
    idf: dict,
    k1: float = 1.2,
    b: float = 0.75,
) -> float:
    """轻量 BM25 计分：预计算词频 + 全局 idf + 长度归一，仍零 embedding。

    idf 由调用方在绑定文档全量切片上一次算好传入。
    """
    if not query_tokens or not tf or not idf or avg_len <= 0:
        return 0.0
    score = 0.0
    unique_q = set(query_tokens)
    norm = k1 * (1.0 - b + b * (doc_len / avg_len))
    for t in unique_q:
        c = tf.get(t, 0)
        if c <= 0:
            continue
        idf_t = idf.get(t, 0.0)
        if idf_t <= 0:
            continue
        # 单字非 ASCII（中文/假名/谚文等）降权，英文整词保持全重
        w = 0.5 if len(t) == 1 and not t.isascii() else 1.0
        score += w * idf_t * (c * (k1 + 1.0)) / (c + norm)
    hits = sum(1 for t in unique_q if tf.get(t, 0) > 0)
    score *= 1.0 + 0.2 * (hits / max(1, len(unique_q)))
    return round(score, 4)


def strip_html(raw: str) -> str:
    """移除 HTML 标签及内嵌脚本（块级标签先换行，避免正文粘连）。"""
    text = re.sub(r"<(script|style).*?</\1>", " ", raw, flags=re.S | re.I)
    text = re.sub(r"</(p|div|br|li|tr|h[1-6]|section|article|header|footer)[^>]*>", "\n", text, flags=re.I)
    text = re.sub(r"<[^>]+>", " ", text)
    return re.sub(r"\n{3,}", "\n\n", re.sub(r"[ \t\xa0]+", " ", text)).strip()


def extract_text_from_bytes(suffix: str, data: bytes) -> str:
    """按文件后缀提取纯文本。"""
    suffix = (suffix or "").lower()

    if suffix in (".png", ".jpg", ".jpeg", ".webp", ".gif", ".bmp"):
        raise RuntimeError(f"不支持的图片类型 {suffix}，请上传文档类文件。")

    if suffix in (".md", ".markdown", ".txt", ".json", ".csv", ".log",
                  ".yaml", ".yml", ".toml", ".ini", ".cfg"):
        return _cap_extract(data.decode("utf-8", errors="ignore"))
    if suffix in (".html", ".htm"):
        return _cap_extract(strip_html(data.decode("utf-8", errors="ignore")))
    if suffix == ".pdf":
        try:
            from pypdf import PdfReader
        except ImportError as e:
            raise RuntimeError("缺少 pypdf 依赖，请 pip install pypdf 后重试。") from e
        reader = PdfReader(io.BytesIO(data))
        texts = []
        for p in reader.pages:
            try:
                t = p.extract_text()
            except Exception:
                continue
            if t and t.strip():
                texts.append(t.strip())
        return _cap_extract("\n\n".join(texts))
    if suffix == ".docx":
        try:
            import docx
        except ImportError as e:
            raise RuntimeError("缺少 python-docx 依赖，请 pip install python-docx 后重试。") from e
        doc = docx.Document(io.BytesIO(data))
        # 按文档块顺序交错提取（段落在表格中间时不再被搬到文末）；解析失败再回退旧逻辑
        try:
            from docx.oxml.table import CT_Tbl
            from docx.oxml.text.paragraph import CT_P
            from docx.table import Table
            from docx.text.paragraph import Paragraph
            parts = []
            for child in doc.element.body.iterchildren():
                if isinstance(child, CT_P):
                    t = Paragraph(child, doc).text.strip()
                    if t:
                        parts.append(t)
                elif isinstance(child, CT_Tbl):
                    for row in Table(child, doc).rows:
                        for cell in row.cells:
                            t = (cell.text or "").strip()
                            if t:
                                parts.append(t)
            if parts:
                return _cap_extract("\n".join(parts))
        except Exception:
            pass
        parts = [p.text.strip() for p in doc.paragraphs if p.text and p.text.strip()]
        for table in doc.tables:  # 回退：旧逻辑（顺序丢失，仅保底）
            for row in table.rows:
                for cell in row.cells:
                    t = (cell.text or "").strip()
                    if t:
                        parts.append(t)
        return _cap_extract("\n".join(parts))
    # 其他兜底当文本解码
    text = data.decode("utf-8", errors="ignore")
    if not text.strip():
        raise RuntimeError(f"不支持的文件类型 {suffix} 或内容无法解码")
    return _cap_extract(text)
